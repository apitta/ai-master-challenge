"""O score da aplicação precisa bater 100% com o calculado pela planilha."""

from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.config import get_settings
from app.scoring import excel_round, score_opportunity
from app.seed import read_workbook


@pytest.fixture(scope="module")
def workbook():
    return read_workbook(get_settings().seed_file)


@pytest.fixture(scope="module")
def excel_results():
    """Valores calculados e salvos pelo Excel (cache das fórmulas), por ID."""
    ws = load_workbook(get_settings().seed_file, data_only=True)["Pipeline"]
    return {
        r[0]: dict(value_points=r[10], time_points=r[11], multiplier=r[12], score=r[13], band=r[15])
        for r in ws.iter_rows(min_row=2, values_only=True)
        if r[0]
    }


def test_params_read_from_sheet(workbook):
    p = workbook["params"]
    assert p.stall_limit_days == 138
    assert [b.start for b in p.multiplier_bands] == [0, 139, 181, 271]
    assert [b.start for b in p.time_bands] == [0, 15, 91]
    assert {r.band: r.min_score for r in p.band_rules} == {"A": 75, "B": 55, "C": 35, "D": 0, "R": None}


def test_every_opportunity_matches_excel(workbook, excel_results):
    params = workbook["params"]
    mismatches = []
    for o in workbook["opportunities"]:
        got = score_opportunity(o["stage"], o["engage_date"], o["product_name"], params)
        expected = excel_results[o["id"]]
        actual = dict(value_points=got.value_points, time_points=got.time_points,
                      multiplier=got.multiplier, score=got.score, band=got.band)
        if actual != expected:
            mismatches.append((o["id"], actual, expected))
    assert len(workbook["opportunities"]) == 2089
    assert not mismatches, mismatches[:5]


def test_excel_round_half_away_from_zero():
    assert excel_round(14.5) == 15  # round(14.5) no Python daria 14
    assert excel_round(22.000000000000004) == 22


@pytest.mark.parametrize("days, band_expected, mult", [(138, "A", 1.0), (139, "R", 0.6), (181, "R", 0.4), (271, "R", 0.25)])
def test_stall_boundaries(workbook, days, band_expected, mult):
    params = workbook["params"]
    engage = params.reference_date - timedelta(days=days)
    result = score_opportunity("Engaging", engage, "GTK 500", params)
    assert result.band == band_expected
    assert result.multiplier == mult


def test_prospecting_gets_fixed_time_points(workbook):
    result = score_opportunity("Prospecting", None, "MG Special", workbook["params"])
    assert (result.value_points, result.time_points, result.score, result.band) == (0, 10, 10, "D")
